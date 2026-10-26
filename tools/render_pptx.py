#!/usr/bin/env python3
"""render_pptx - render an output-contract deck payload to a .pptx file.

Optional handoff step for the One-Click AI PPT skill. Reads the structured
payload described in references/output-contract.md and writes one PowerPoint
slide per payload slide, into the bundled Office master or a corporate template
supplied with --template.

A slide may carry a structured `chart` or `table` block next to its bullets. Those
render as native PowerPoint objects (editable data, not pictures), and a slide is
laid out in fixed layers: headline band, copy zone, data zone, source footnote.
Token-only `visual` text keeps the old behaviour and lands in the speaker notes.

    python tools/render_pptx.py examples/status-report.json
    python tools/render_pptx.py deck.json --template acme.potx --master-map acme-master-map.json
    python tools/render_pptx.py deck.json --check
    python tools/render_pptx.py deck.json --out out/acme_deck_v1.0.0_20260922.pptx

Conventions:
    * layout + placeholder resolution - references/master-mapping.md
    * output file naming              - references/export-checklist.md
    * payload schema                  - references/output-contract.md

Requires python-pptx (pip install python-pptx); the deck checker
tools/deck_lint.py stays dependency-free. Exit codes: 0 ok, 1 hard failure,
2 missing dependency.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

EXIT_OK = 0
EXIT_FAIL = 1
EXIT_MISSING_DEP = 2

SCHEMA_VERSIONS = ("1.0", "1.1")
SCHEMA_VERSION = SCHEMA_VERSIONS[-1]
ROLES = ("title", "content", "section", "closing")

# Layered layout, in inches, for the default 16:9 body. A slide is built bottom-up
# from these bands; nothing may share a band with another layer.
LAYERS = ("headline", "copy", "data", "footnote")
HEADLINE_BAND = (0.55, 0.35, 12.23, 1.05)      # x, y, w, h
BODY_TOP = 1.62
BODY_BOTTOM = 6.30
LEFT_GUTTER = 0.62
COLUMN_GAP = 0.42
DATA_SHARE = 0.56                               # data zone share of the body width
FOOTNOTE_BAND = (0.62, 6.42, 12.23, 0.38)

CHART_LIMITS = {"categories": (2, 12), "series": (1, 4), "points": 12}
TABLE_LIMITS = {"columns": (2, 6), "rows": (1, 8)}
CHART_TYPE_ALIASES = {
    "column": "COLUMN_CLUSTERED",
    "bar": "BAR_CLUSTERED",
    "line": "LINE_MARKERS",
    "area": "AREA",
    "pie": "PIE",
    "doughnut": "DOUGHNUT",
    "stacked-column": "COLUMN_STACKED",
    "stacked-bar": "BAR_STACKED",
}
CHART_TYPE_NAMES = tuple(CHART_TYPE_ALIASES)
PIE_TYPES = ("pie", "doughnut")
DEFAULT_LAYOUT_INDEX = {"title": 0, "content": 1}
DEFAULT_VERSION = "v1.0.0"

SLUG_RE = re.compile(r"[^a-z0-9]+")


class RenderError(Exception):
    """Hard failure: bad payload, unreadable template or mapping."""


# ------------------------------------------------------------------- utilities


def ensure_utf8_stdout() -> None:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8", errors="replace")
            except (ValueError, OSError):
                pass


def slugify(text: str, fallback: str = "deck") -> str:
    slug = SLUG_RE.sub("-", text.strip().lower()).strip("-")
    return slug or fallback


def load_json(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as err:
        raise RenderError(f"file not found: {path}") from err
    except json.JSONDecodeError as err:
        raise RenderError(f"invalid JSON in {path}: {err}") from err


def import_pptx():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import PP_PLACEHOLDER
        from pptx.util import Inches
    except ImportError:
        print(
            "python-pptx is not installed; rendering skipped.\n"
            "    pip install python-pptx\n"
            "The no-dependency path is still available: hand over the Markdown "
            "deck and the JSON payload.",
            file=sys.stderr,
        )
        raise SystemExit(EXIT_MISSING_DEP)
    return Presentation, PP_PLACEHOLDER, Inches


# -------------------------------------------------------------------- payload


def validate_payload(data: dict) -> list[str]:
    """Return hard errors; the payload must be clean before rendering."""
    errors: list[str] = []
    if not isinstance(data, dict):
        return ["payload root must be an object"]

    meta = data.get("meta")
    slides = data.get("slides")
    if not isinstance(meta, dict):
        errors.append("meta object is missing")
        meta = {}
    if not isinstance(slides, list) or not slides:
        errors.append("slides array is missing or empty")
        slides = []

    title = meta.get("title")
    if not isinstance(title, str) or not title.strip():
        errors.append("meta.title is empty")

    declared = meta.get("slide_count")
    if declared is not None and declared != len(slides):
        errors.append(f"meta.slide_count ({declared}) != len(slides) ({len(slides)})")

    schema = data.get("schema_version")
    if schema is not None and str(schema) not in SCHEMA_VERSIONS:
        errors.append(f"unsupported schema_version {schema!r} (expected one of {SCHEMA_VERSIONS})")

    for index, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            errors.append(f"slide {index} is not an object")
            continue
        headline = slide.get("title")
        if not isinstance(headline, str) or not headline.strip():
            errors.append(f"slide {index} has no title")
        bullets = slide.get("bullets")
        if not isinstance(bullets, list) or not bullets:
            errors.append(f"slide {index} has no bullets")
        elif not all(isinstance(item, str) and item.strip() for item in bullets):
            errors.append(f"slide {index} has an empty bullet")
        # Optional data layer: chart or table, never both, and never on the title slide.
        has_chart = slide.get("chart") is not None
        has_table = slide.get("table") is not None
        if has_chart and has_table:
            errors.append(f"slide {index} carries both chart and table; pick one data layer")
        if (has_chart or has_table) and index == 1:
            errors.append(f"slide {index} is the title slide; move the data block to a content slide")
        if has_chart:
            note_errors, _ = check_chart(slide.get("chart"), f"slide {index}.chart")
            errors.extend(note_errors)
        if has_table:
            note_errors, _ = check_table(slide.get("table"), f"slide {index}.table")
            errors.extend(note_errors)
    return errors


# ----------------------------------------------------------------- data layer


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def check_chart(raw, where: str) -> tuple[list[str], list[str]]:
    """Validate a chart block; returns (errors, warnings) without mutating input."""
    errors: list[str] = []
    warnings: list[str] = []
    if not isinstance(raw, dict):
        return [f"{where} must be an object"], warnings

    kind = raw.get("type", "column")
    kind = str(kind).strip().lower()
    if kind not in CHART_TYPE_ALIASES:
        errors.append(f"{where}.type must be one of {', '.join(CHART_TYPE_NAMES)}; got {kind!r}")

    low, high = CHART_LIMITS["categories"]
    categories = raw.get("categories")
    if not isinstance(categories, list) or not low <= len(categories) <= high:
        errors.append(f"{where}.categories needs {low}-{high} items")
        categories = []
    elif not all(isinstance(item, (str, int, float)) and str(item).strip() for item in categories):
        errors.append(f"{where}.categories has an empty label")
        categories = []

    series = raw.get("series")
    low_s, high_s = CHART_LIMITS["series"]
    if not isinstance(series, list) or not low_s <= len(series) <= high_s:
        errors.append(f"{where}.series needs {low_s}-{high_s} items")
        series = []
    else:
        width = len(categories)
        for pos, entry in enumerate(series, start=1):
            if not isinstance(entry, dict):
                errors.append(f"{where}.series[{pos}] must be an object")
                continue
            values = entry.get("values")
            if not isinstance(values, list) or len(values) != width:
                errors.append(f"{where}.series[{pos}].values must have {width} numbers")
                continue
            if not all(_is_number(item) for item in values):
                errors.append(f"{where}.series[{pos}].values must be numeric")

    emphasis = raw.get("emphasis")
    if emphasis is not None:
        if not isinstance(emphasis, list) or not all(isinstance(i, int) for i in emphasis):
            errors.append(f"{where}.emphasis must be a list of category indexes")
        elif categories and any(i < 0 or i >= len(categories) for i in emphasis):
            errors.append(f"{where}.emphasis has an index outside the categories")

    if kind in PIE_TYPES and isinstance(series, list) and len(series) > 1:
        warnings.append(f"{where}: {kind} charts plot one series, extra series ignored")
    return errors, warnings


def check_table(raw, where: str) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    if not isinstance(raw, dict):
        return [f"{where} must be an object"], warnings

    low, high = TABLE_LIMITS["columns"]
    columns = raw.get("columns")
    if not isinstance(columns, list) or not low <= len(columns) <= high:
        errors.append(f"{where}.columns needs {low}-{high} headers")
        columns = []
    elif not all(isinstance(item, str) and item.strip() for item in columns):
        errors.append(f"{where}.columns has an empty header")
        columns = []

    low_r, high_r = TABLE_LIMITS["rows"]
    rows = raw.get("rows")
    if not isinstance(rows, list) or not low_r <= len(rows) <= high_r:
        errors.append(f"{where}.rows needs {low_r}-{high_r} rows")
        rows = []
    else:
        for pos, row in enumerate(rows, start=1):
            if not isinstance(row, list) or len(row) != len(columns):
                errors.append(f"{where}.rows[{pos}] must have {len(columns)} cells")
            elif not all(isinstance(cell, (str, int, float)) for cell in row):
                errors.append(f"{where}.rows[{pos}] has a non-scalar cell")
    return errors, warnings


def plan_layers(slide_w: float, slide_h: float, has_data: bool) -> dict:
    """Split the body into fixed bands; the data band never overlaps the copy zone."""
    body_w = slide_w - 2 * LEFT_GUTTER
    body_h = min(BODY_BOTTOM, slide_h - 0.55) - BODY_TOP
    copy = (LEFT_GUTTER, BODY_TOP, body_w, body_h)
    data = None
    note = (FOOTNOTE_BAND[0], slide_h - 1.05, body_w, FOOTNOTE_BAND[3])
    if has_data:
        copy_w = round(body_w * (1 - DATA_SHARE) - COLUMN_GAP / 2, 3)
        data_w = round(body_w - copy_w - COLUMN_GAP, 3)
        copy = (LEFT_GUTTER, BODY_TOP, copy_w, body_h)
        data = (round(LEFT_GUTTER + copy_w + COLUMN_GAP, 3), BODY_TOP, data_w, body_h)
    headline = (HEADLINE_BAND[0], HEADLINE_BAND[1], round(slide_w - 2 * HEADLINE_BAND[0], 3),
                HEADLINE_BAND[3])
    return {"headline": headline, "copy": copy, "data": data, "footnote": note,
            "slide": (slide_w, slide_h)}


def _place(shape, zone, Inches) -> None:
    x, y, w, h = zone
    shape.left, shape.top = Inches(x), Inches(y)
    shape.width, shape.height = Inches(w), Inches(h)


def draw_chart(slide, chart_spec: dict, zone, where: str, warnings: list[str]) -> str:
    """Add a native chart; returns the chart family actually used."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    kind = str(chart_spec.get("type", "column")).strip().lower()
    chart_data = CategoryChartData()
    chart_data.categories = [str(item) for item in chart_spec["categories"]]
    series = chart_spec["series"]
    if kind in PIE_TYPES:
        series = series[:1]
    for entry in series:
        chart_data.add_series(str(entry.get("name") or "Series"), entry["values"])

    chart_type = getattr(XL_CHART_TYPE, CHART_TYPE_ALIASES[kind])
    frame = slide.shapes.add_chart(chart_type, 0, 0, 1, 1, chart_data)
    _place(frame, zone, Inches)
    chart = frame.chart

    title = chart_spec.get("title")
    if isinstance(title, str) and title.strip():
        chart.has_title = True
        chart.chart_title.text_frame.text = title.strip()
        chart.chart_title.text_frame.paragraphs[0].font.size = Pt(14)
        chart.chart_title.text_frame.paragraphs[0].font.bold = True
    else:
        chart.has_title = False

    chart.has_legend = len(series) > 1 and kind not in PIE_TYPES
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)

    unit = chart_spec.get("unit")
    if isinstance(unit, str) and unit.strip() and kind not in PIE_TYPES:
        try:
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = unit.strip()
            chart.value_axis.axis_title.text_frame.paragraphs[0].font.size = Pt(11)
        except (AttributeError, ValueError) as err:  # pragma: no cover - chart dependent
            warnings.append(f"{where}: could not label the value axis ({err})")

    try:
        chart.font.size = Pt(12)
        chart.font.name = "Segoe UI"
    except (AttributeError, ValueError):  # pragma: no cover - chart dependent
        pass

    emphasis = chart_spec.get("emphasis") or []
    if emphasis:
        try:
            points = chart.plots[0].series[0].points
            for index in range(len(chart_spec["categories"])):
                fill = points[index].format.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(0x1F, 0x38, 0x64) if index in emphasis else RGBColor(0xBD, 0xD0, 0xE3)
        except (AttributeError, IndexError):  # pragma: no cover - chart dependent
            warnings.append(f"{where}: emphasis colours not supported by this chart type")

    return kind


def draw_table(slide, table_spec: dict, zone, where: str, warnings: list[str]) -> dict:
    """Add a native table; returns the geometry actually written."""
    from pptx.util import Inches, Pt

    columns = [str(item) for item in table_spec["columns"]]
    rows = [[str(cell) for cell in row] for row in table_spec["rows"]]
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), 0, 0, 1, 1)
    _place(shape, zone, Inches)
    table = shape.table

    for pos, header in enumerate(columns):
        cell = table.cell(0, pos)
        cell.text = header
        cell.text_frame.paragraphs[0].font.size = Pt(12)
        cell.text_frame.paragraphs[0].font.bold = True
    for row_pos, row in enumerate(rows, start=1):
        for col_pos, value in enumerate(row):
            cell = table.cell(row_pos, col_pos)
            cell.text = value
            cell.text_frame.paragraphs[0].font.size = Pt(12)
    return {"columns": len(columns), "rows": len(rows)}


# ------------------------------------------------------------ master mapping


def resolve_layout_config(mapping: dict | None) -> dict:
    layouts = {}
    if isinstance(mapping, dict):
        raw = mapping.get("layouts")
        if isinstance(raw, dict):
            layouts = {role: cfg for role, cfg in raw.items() if isinstance(cfg, dict)}
    return layouts


def match_layout(prs, role: str, config: dict, report: list[dict]):
    """Resolve one role to a slide layout: name, keywords, index, then default."""
    names = [layout.name for layout in prs.slide_layouts]

    wanted_name = config.get("name")
    if isinstance(wanted_name, str) and wanted_name.strip():
        for layout in prs.slide_layouts:
            if layout.name.strip().lower() == wanted_name.strip().lower():
                report.append({"role": role, "layout": layout.name, "via": "name"})
                return layout

    keywords = config.get("keywords")
    if isinstance(keywords, list):
        for keyword in keywords:
            if not isinstance(keyword, str) or not keyword.strip():
                continue
            needle = keyword.strip().lower()
            for layout in prs.slide_layouts:
                if needle in layout.name.lower():
                    report.append({"role": role, "layout": layout.name, "via": f"keyword:{keyword}"})
                    return layout

    index = config.get("index")
    if isinstance(index, int) and 0 <= index < len(prs.slide_layouts):
        layout = prs.slide_layouts[index]
        report.append({"role": role, "layout": layout.name, "via": "index"})
        return layout

    fallback = DEFAULT_LAYOUT_INDEX.get(role, 1)
    if fallback < len(prs.slide_layouts):
        layout = prs.slide_layouts[fallback]
        report.append({"role": role, "layout": layout.name, "via": "default"})
        return layout

    raise RenderError(f"no usable layout for role {role!r} (template has {len(names)} layouts)")


def pick_placeholder(container, config: dict, kind: str, placeholder_types, warnings: list[str]):
    """Resolve a placeholder by explicit idx, then by type, then give up.

    ``container`` is the *slide* (layout placeholders are only the blueprint;
    add_slide clones them onto the slide and only the slide copies are visible).
    """
    wanted_idx = None
    placeholders_cfg = config.get("placeholders")
    if isinstance(placeholders_cfg, dict):
        candidate = placeholders_cfg.get(kind)
        if isinstance(candidate, int):
            wanted_idx = candidate

    if wanted_idx is not None:
        for placeholder in container.placeholders:
            if placeholder.placeholder_format.idx == wanted_idx:
                return placeholder
        warnings.append(f"placeholder idx {wanted_idx} for {kind!r} not found on the slide")

    for placeholder in container.placeholders:
        if placeholder.placeholder_format.type in placeholder_types:
            return placeholder
    return None


# ------------------------------------------------------------------- rendering


def assign_roles(count: int) -> list[str]:
    roles = ["content"] * count
    if count:
        roles[0] = "title"
    return roles


def write_bullets(text_frame, bullets: list[str]) -> None:
    text_frame.text = bullets[0]
    for bullet in bullets[1:]:
        paragraph = text_frame.add_paragraph()
        paragraph.text = bullet


def set_notes(slide, notes: str, visual: str | None, warnings: list[str]) -> None:
    body = notes.strip()
    if visual and visual.strip():
        body = f"{body}\n\nVisual: {visual.strip()}"
    try:
        slide.notes_slide.notes_text_frame.text = body
    except (AttributeError, ValueError) as err:  # pragma: no cover - template dependent
        warnings.append(f"could not write speaker notes: {err}")


def render_deck(payload: dict, template: Path | None, mapping: dict | None, check_only: bool,
                out_path: Path | None, force: bool) -> dict:
    meta = payload["meta"]
    slides = payload["slides"]
    assignment = assign_roles(len(slides))

    if check_only and template is None:
        layout_config = resolve_layout_config(mapping)
        return {
            "mode": "check",
            "payload_slides": len(slides),
            "template": None,
            "template_layouts": [],
            "layout_resolution": [
                {
                    "role": role,
                    "layout": f"built-in default index {DEFAULT_LAYOUT_INDEX.get(role, 1)}",
                    "via": "default",
                    "configured": bool(layout_config.get(role)),
                }
                for role in ROLES
            ],
            "per_slide": [{"slide": index, "role": role} for index, role in enumerate(assignment, start=1)],
            "mapping_applied": bool(mapping),
            "note": "payload validated; python-pptx is only needed with --template",
        }

    Presentation, PP_PLACEHOLDER, Inches = import_pptx()
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    try:
        prs = Presentation(str(template)) if template else Presentation()
    except Exception as err:  # pragma: no cover - template dependent
        raise RenderError(f"cannot open template {template}: {err}") from err

    title_types = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    subtitle_types = (PP_PLACEHOLDER.SUBTITLE,)
    body_types = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)

    layout_config = resolve_layout_config(mapping)
    layout_report: list[dict] = []
    resolved = {role: match_layout(prs, role, layout_config.get(role, {}), layout_report)
                for role in ROLES}

    if check_only:
        per_slide = []
        for index, role in enumerate(assignment, start=1):
            layout = resolved[role]
            per_slide.append({"slide": index, "role": role, "layout": layout.name})
        return {
            "mode": "check",
            "payload_slides": len(slides),
            "template": str(template) if template else None,
            "template_layouts": [layout.name for layout in prs.slide_layouts],
            "layout_resolution": layout_report,
            "per_slide": per_slide,
            "mapping_applied": bool(mapping),
        }

    warnings: list[str] = []
    if template is not None or mapping is not None:
        # Only meaningful when the caller asked for a specific master.
        warnings = [
            f"layout fallback for role {entry['role']!r} (template layout names did not match)"
            for entry in layout_report
            if entry["via"] == "default"
        ]

    built: list[dict] = []
    for index, (role, slide_data) in enumerate(zip(assignment, slides), start=1):
        layout = resolved[role]
        slide = prs.slides.add_slide(layout)
        config = layout_config.get(role, {})

        title_ph = pick_placeholder(slide, config, "title", title_types, warnings)
        if title_ph is not None:
            if role == "title":
                title_ph.text_frame.text = str(meta.get("title", "")).strip()
            else:
                title_ph.text_frame.text = str(slide_data.get("title", "")).strip()
        else:
            warnings.append(f"slide {index}: no title placeholder in layout {layout.name!r}")

        chart_spec = slide_data.get("chart")
        table_spec = slide_data.get("table")
        has_data = isinstance(chart_spec, dict) or isinstance(table_spec, dict)
        zones = plan_layers(prs.slide_width / 914400, prs.slide_height / 914400, has_data)
        layers = {"headline": True, "copy": 0, "data": None, "footnote": None}

        if role == "title":
            subtitle = f"{meta.get('audience', '')} · {meta.get('purpose', '')} · ~{meta.get('est_minutes', '')} min"
            subtitle_ph = pick_placeholder(slide, config, "subtitle", subtitle_types, warnings)
            if subtitle_ph is not None:
                subtitle_ph.text_frame.text = subtitle.strip(" ·")
        else:
            bullets = [str(item) for item in slide_data.get("bullets", [])]
            if has_data and len(bullets) > 3:
                warnings.append(f"slide {index}: {len(bullets)} bullets beside a data block, kept 3")
                bullets = bullets[:3]
            body_ph = pick_placeholder(slide, config, "body", body_types, warnings)
            if body_ph is not None:
                if has_data:
                    _place(body_ph, zones["copy"], Inches)
                write_bullets(body_ph.text_frame, bullets)
            else:
                box = slide.shapes.add_textbox(0, 0, 1, 1)
                _place(box, zones["copy"], Inches)
                write_bullets(box.text_frame, bullets)
                warnings.append(f"slide {index}: no body placeholder, added a text box")
            layers["copy"] = len(bullets)

            if isinstance(chart_spec, dict):
                _, note_warnings = check_chart(chart_spec, f"slide {index}.chart")
                warnings.extend(note_warnings)
                try:
                    used = draw_chart(slide, chart_spec, zones["data"], f"slide {index}.chart", warnings)
                    layers["data"] = f"chart:{used}"
                    source = chart_spec.get("source")
                except Exception as err:  # pragma: no cover - chart dependent
                    warnings.append(f"slide {index}: chart dropped ({err})")
                    source = None
            elif isinstance(table_spec, dict):
                source = table_spec.get("source")
                try:
                    info = draw_table(slide, table_spec, zones["data"], f"slide {index}.table", warnings)
                    layers["data"] = f"table:{info['columns']}x{info['rows']}"
                except Exception as err:  # pragma: no cover - table dependent
                    warnings.append(f"slide {index}: table dropped ({err})")
            else:
                source = None

            caption = source if isinstance(source, str) and source.strip() else None
            if caption:
                note_h = zones["footnote"]
                note_frame = slide.shapes.add_textbox(
                    Inches(note_h[0]), Inches(note_h[1]), Inches(note_h[2]), Inches(note_h[3]))
                para = note_frame.text_frame.paragraphs[0]
                para.text = f"Source: {caption.strip()}"
                para.font.size = Pt(10)
                para.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
                layers["footnote"] = caption.strip()

        set_notes(slide, str(slide_data.get("speaker_notes", "")), slide_data.get("visual"), warnings)
        built.append({"slide": index, "role": role, "layout": layout.name,
                      "bullets": layers["copy"], "layers": layers})

    if out_path is None:
        raise RenderError("output path missing")
    if out_path.exists() and not force:
        raise RenderError(f"{out_path} already exists; pass --force to overwrite")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))

    return {
        "mode": "render",
        "output": str(out_path),
        "payload_slides": len(slides),
        "template": str(template) if template else None,
        "mapping_applied": bool(mapping),
        "layout_resolution": layout_report,
        "slides": built,
        "warnings": warnings,
    }


def default_out_path(payload_path: Path, meta: dict, version: str, slot: int | None) -> Path:
    slug = slugify(str(meta.get("title", payload_path.stem)))
    stamp = date.today().strftime("%Y%m%d")
    name = f"{slug}_{version}_{slot}min_{stamp}.pptx" if slot else f"{slug}_{version}_{stamp}.pptx"
    return payload_path.parent / name


# ------------------------------------------------------------------------ cli


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Render an output-contract deck payload (.json) to a .pptx file.",
    )
    parser.add_argument("payload", help="deck payload JSON (see references/output-contract.md)")
    parser.add_argument("--template", help="corporate .pptx/.potx template supplying the masters")
    parser.add_argument("--master-map", help="layout + placeholder mapping JSON (references/master-mapping.md)")
    parser.add_argument("--out", help="output .pptx path (default: naming convention in export-checklist.md)")
    parser.add_argument("--version", default=DEFAULT_VERSION, help=f"version token for the file name (default: {DEFAULT_VERSION})")
    parser.add_argument("--slot", type=int, help="duration slot in minutes for scaled packs (3 / 8 / 15)")
    parser.add_argument("--check", action="store_true", help="validate payload and layout mapping, write nothing")
    parser.add_argument("--force", action="store_true", help="overwrite an existing output file")
    parser.add_argument("--json", action="store_true", help="print the report as JSON")
    return parser


def main(argv: list[str] | None = None) -> int:
    ensure_utf8_stdout()
    args = build_parser().parse_args(argv)

    try:
        payload_path = Path(args.payload)
        payload = load_json(payload_path)
        errors = validate_payload(payload)
        if errors:
            raise RenderError("payload invalid:\n  - " + "\n  - ".join(errors))

        mapping = None
        if args.master_map:
            mapping = load_json(Path(args.master_map))
            if not isinstance(mapping, dict) or not isinstance(mapping.get("layouts"), dict):
                raise RenderError(f"master map {args.master_map} must contain a 'layouts' object")
        elif args.template is None and not args.check:
            mapping = None

        template = Path(args.template) if args.template else None
        if template is not None and not template.exists():
            raise RenderError(f"template not found: {template}")
        if args.master_map and template is None:
            print("note: --master-map without --template; the bundled master and its defaults are used",
                  file=sys.stderr)

        out_path = Path(args.out) if args.out else None
        if out_path is None and not args.check:
            out_path = default_out_path(payload_path, payload.get("meta", {}), args.version, args.slot)

        report = render_deck(payload, template, mapping, args.check, out_path, args.force)
    except RenderError as err:
        print(f"ERROR {err}", file=sys.stderr)
        return EXIT_FAIL

    if args.json:
        print(json.dumps(report, indent=2, ensure_ascii=False))
    else:
        if report["mode"] == "check":
            print(f"check ok: {report['payload_slides']} slides, template={report['template'] or 'bundled'}")
            for entry in report["layout_resolution"]:
                print(f"  {entry['role']:<8} -> {entry['layout']}  ({entry['via']})")
        else:
            print(f"rendered {report['payload_slides']} slides -> {report['output']}")
            for entry in report["layout_resolution"]:
                print(f"  {entry['role']:<8} -> {entry['layout']}  ({entry['via']})")
            for entry in report["slides"]:
                if entry["layers"]["data"]:
                    print(f"  slide {entry['slide']}: data layer {entry['layers']['data']}")
            for warning in report["warnings"]:
                print(f"  warning: {warning}")
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
