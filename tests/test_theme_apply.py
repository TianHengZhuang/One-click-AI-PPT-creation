#!/usr/bin/env python3
"""Tests for tools/theme_apply.py - colour maths, theme validation, restyling.

The restyle tests render a real deck and are skipped when python-pptx is absent.

Run with:
    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

import theme_apply as ta  # noqa: E402

HAVE_PPTX = importlib.util.find_spec("pptx") is not None
THEME_DIR = ROOT / "assets" / "themes"


def theme_copy(name: str = "swiss") -> dict:
    return json.loads((THEME_DIR / f"{name}.json").read_text(encoding="utf-8"))


class ColourMathTests(unittest.TestCase):
    def test_hex_parsing(self):
        self.assertEqual(ta._hex_to_rgb("#1A2B3C"), (26, 43, 60))
        self.assertEqual(ta._hex_to_rgb("1a2b3c"), (26, 43, 60))

    def test_short_hex_rejected(self):
        with self.assertRaises(ValueError):
            ta._hex_to_rgb("#FFF")

    def test_non_hex_rejected(self):
        with self.assertRaises(ValueError):
            ta._hex_to_rgb("#GGGGGG")

    def test_luminance_bounds(self):
        self.assertAlmostEqual(ta.relative_luminance("#FFFFFF"), 1.0, places=3)
        self.assertAlmostEqual(ta.relative_luminance("#000000"), 0.0, places=3)

    def test_black_on_white_is_21(self):
        self.assertEqual(ta.contrast_ratio("#000000", "#FFFFFF"), 21.0)

    def test_contrast_is_symmetric(self):
        self.assertEqual(
            ta.contrast_ratio("#123456", "#FDFDFD"),
            ta.contrast_ratio("#FDFDFD", "#123456"),
        )

    def test_identical_colours_are_1(self):
        self.assertEqual(ta.contrast_ratio("#445566", "#445566"), 1.0)


class CheckThemeTests(unittest.TestCase):
    def test_bundled_themes_are_valid(self):
        for name in ("consulting", "swiss", "editorial"):
            errors = ta.check_theme(theme_copy(name), where=name)
            self.assertEqual(errors, [], msg=f"{name}: {errors}")

    def test_missing_required_key(self):
        theme = theme_copy()
        del theme["font"]
        self.assertTrue(any("font" in err for err in ta.check_theme(theme)))

    def test_missing_palette_key(self):
        theme = theme_copy()
        del theme["palette"]["accent"]
        self.assertTrue(any("palette.accent" in err for err in ta.check_theme(theme)))

    def test_bad_hex(self):
        theme = theme_copy()
        theme["palette"]["ink"] = "navy"
        self.assertTrue(any("palette.ink" in err for err in ta.check_theme(theme)))

    def test_low_ink_contrast_rejected(self):
        theme = theme_copy()
        theme["palette"]["ink"] = "#CCCCCC"
        errors = ta.check_theme(theme)
        self.assertTrue(any("unreadable" in err for err in errors))

    def test_low_accent_contrast_rejected(self):
        theme = theme_copy()
        theme["palette"]["accent"] = "#F5F0E6"
        self.assertTrue(any("accent" in err and "below" in err for err in ta.check_theme(theme)))

    def test_unreadable_text_on_accent(self):
        theme = theme_copy()
        theme["palette"]["on_accent"] = theme["palette"]["accent"]
        self.assertTrue(any("on_accent" in err for err in ta.check_theme(theme)))

    def test_font_size_floor(self):
        theme = theme_copy()
        theme["size"]["body"] = 7
        self.assertTrue(any("9" in err and "body" in err for err in ta.check_theme(theme)))

    def test_title_must_outrank_body(self):
        theme = theme_copy()
        theme["size"]["title"] = theme["size"]["body"]
        self.assertTrue(any("title size" in err for err in ta.check_theme(theme)))

    def test_chart_series_must_be_a_list(self):
        theme = theme_copy()
        theme["chart_series"] = []
        self.assertTrue(any("chart_series" in err for err in ta.check_theme(theme)))

    def test_chart_series_colour_must_be_hex(self):
        theme = theme_copy()
        theme["chart_series"] = ["#111111", "red"]
        self.assertTrue(any("chart_series[2]" in err for err in ta.check_theme(theme)))

    def test_chart_series_is_optional(self):
        theme = theme_copy()
        theme.pop("chart_series", None)
        self.assertEqual(ta.check_theme(theme), [])


class ResolveThemeTests(unittest.TestCase):
    def test_resolve_by_name(self):
        theme, source = ta.resolve_theme("swiss", THEME_DIR)
        self.assertEqual(theme["name"], "swiss")
        self.assertTrue(source.endswith("swiss.json"))

    def test_resolve_by_path(self):
        theme, source = ta.resolve_theme(str(THEME_DIR / "editorial.json"))
        self.assertEqual(theme["name"], "editorial")
        self.assertTrue(source.endswith("editorial.json"))

    def test_unknown_name_lists_bundled(self):
        with self.assertRaises(ta.ThemeError) as ctx:
            ta.resolve_theme("brutalist", THEME_DIR)
        self.assertIn("swiss", str(ctx.exception))

    def test_missing_path_raises(self):
        with self.assertRaises(ta.ThemeError):
            ta.resolve_theme(str(THEME_DIR / "nope.json"), THEME_DIR)

    def test_invalid_json_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "broken.json"
            path.write_text("{not json", encoding="utf-8")
            with self.assertRaises(ta.ThemeError):
                ta.resolve_theme(str(path))


class CliTests(unittest.TestCase):
    def test_list_exits_zero(self):
        self.assertEqual(ta.main(["--list"]), 0)

    def test_check_exits_zero_for_bundled_theme(self):
        self.assertEqual(ta.main(["x.pptx", "--theme", "swiss", "--check"]), 0)

    def test_missing_args_exit_one(self):
        self.assertEqual(ta.main([]), 1)

    def test_unknown_theme_exits_one(self):
        self.assertEqual(ta.main(["x.pptx", "--theme", "brutalist"]), 1)

    def test_missing_deck_exits_one(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(ta.main([str(Path(tmp) / "nope.pptx"), "--theme", "swiss"]), 1)


@unittest.skipUnless(HAVE_PPTX, "python-pptx is not installed")
class ApplyTests(unittest.TestCase):
    def render(self, tmp: str) -> Path:
        import render_pptx as rp

        out = Path(tmp) / "deck.pptx"
        code = rp.main([str(ROOT / "examples" / "data-layer.example.json"), "--out", str(out), "--force"])
        self.assertEqual(code, 0)
        return out

    def test_restyles_text_charts_and_tables(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            deck = self.render(tmp)
            code = ta.main([str(deck), "--theme", "editorial", "--out", str(Path(tmp) / "themed.pptx")])
            self.assertEqual(code, 0)
            prs = Presentation(str(Path(tmp) / "themed.pptx"))
            report = ta.apply_theme(prs, ta.resolve_theme("editorial", THEME_DIR)[0])
            self.assertGreater(report["restyled"]["runs"], 0)
            # three charts, four series (the stacked plot carries two)
            self.assertEqual(report["restyled"]["charts"], 4)
            self.assertGreater(report["restyled"]["tables"], 0)

    def test_background_follows_the_theme(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            deck = self.render(tmp)
            theme, _ = ta.resolve_theme("consulting", THEME_DIR)
            prs = Presentation(str(deck))
            ta.apply_theme(prs, theme)
            expected = "%02X%02X%02X" % ta._hex_to_rgb(theme["palette"]["bg"])
            for slide in prs.slides:
                self.assertEqual(str(slide.background.fill.fore_color.rgb), expected)

    def test_run_fonts_follow_the_theme(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as tmp:
            deck = self.render(tmp)
            theme, _ = ta.resolve_theme("swiss", THEME_DIR)
            prs = Presentation(str(deck))
            ta.apply_theme(prs, theme)
            slide = prs.slides[1]
            fonts = {
                run.font.name
                for shape in slide.shapes
                if shape.has_text_frame
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
            }
            self.assertTrue(fonts)
            self.assertTrue(fonts <= {theme["font"]["title"], theme["font"]["body"], theme["font"].get("table")})

    def test_output_name_defaults_to_deck_plus_theme(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = self.render(tmp)
            self.assertEqual(ta.main([str(deck), "--theme", "swiss"]), 0)
            self.assertTrue((Path(tmp) / "deck_swiss.pptx").exists())

    def test_existing_output_needs_force(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = self.render(tmp)
            self.assertEqual(ta.main([str(deck), "--theme", "swiss"]), 0)
            self.assertEqual(ta.main([str(deck), "--theme", "swiss"]), 1)
            self.assertEqual(ta.main([str(deck), "--theme", "swiss", "--force"]), 0)


if __name__ == "__main__":
    unittest.main()
